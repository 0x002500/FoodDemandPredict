"""Production XLSX input/output workflow for canteen dish-sales forecasts."""
from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

import numpy as np
from chinese_calendar import is_workday
from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties
from openpyxl.styles import PatternFill

from predict.t0_forecaster import T0DemandPredictor

DATE_COLUMN = "日期"
COVARIATE_COLUMNS = ("当周的第几天", "是否工作日")
FORECAST_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")


@dataclass(frozen=True)
class ForecastWorkbookReport:
    output_path: Path
    forecast_rows: int


def _headers(worksheet) -> dict[str, int]:
    values = [cell.value for cell in worksheet[1]]
    required = {DATE_COLUMN, *COVARIATE_COLUMNS}
    missing = required.difference(values)
    if missing:
        raise ValueError(f"工作表 {worksheet.title} 缺少列: {sorted(missing)}")
    return {value: index + 1 for index, value in enumerate(values) if value is not None}


def _as_date(value: object, sheet_name: str) -> date:
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"工作表 {sheet_name} 的日期列包含非日期值: {value!r}")


def _dish_columns(columns: dict[str, int]) -> dict[str, int]:
    excluded = {DATE_COLUMN, *COVARIATE_COLUMNS}
    dishes = {name: index for name, index in columns.items() if name not in excluded}
    if not dishes:
        raise ValueError("未找到菜品销量列")
    return dishes


def _historical_rows(worksheet, columns: dict[str, int], dishes: dict[str, int]) -> tuple[np.ndarray, np.ndarray, date]:
    rows: list[tuple[date, float, float, list[float]]] = []
    for row in range(2, worksheet.max_row + 1):
        raw_date = worksheet.cell(row, columns[DATE_COLUMN]).value
        if raw_date is None:
            continue
        try:
            current_date = _as_date(raw_date, worksheet.title)
            day_of_week = worksheet.cell(row, columns[COVARIATE_COLUMNS[0]]).value
            workday = worksheet.cell(row, columns[COVARIATE_COLUMNS[1]]).value
            rows.append((current_date, float(current_date.isoweekday() if day_of_week is None else day_of_week), float(int(is_workday(current_date)) if workday is None else workday), [float(worksheet.cell(row, column).value) for column in dishes.values()]))
        except (TypeError, ValueError) as error:
            raise ValueError(f"工作表 {worksheet.title} 第 {row} 行的数据不可用于预测") from error
    rows.sort(key=lambda item: item[0])
    if not rows:
        raise ValueError(f"工作表 {worksheet.title} 没有可用的历史销量")
    histories = np.asarray([item[3] for item in rows], dtype=np.float32).T
    covariates = np.asarray([[item[1], item[2]] for item in rows], dtype=np.float32)
    return histories, covariates, rows[-1][0]


def _future_covariates(last_date: date, days: int) -> tuple[list[date], np.ndarray]:
    dates = [last_date + timedelta(days=offset) for offset in range(1, days + 1)]
    values = np.asarray([[current.isoweekday(), int(is_workday(current))] for current in dates], dtype=np.float32)
    return dates, values


def _append_rows(worksheet, columns: dict[str, int], dishes: dict[str, int], dates: list[date], predicted: np.ndarray) -> None:
    for horizon_index, current_date in enumerate(dates):
        row = worksheet.max_row + 1
        worksheet.cell(row, columns[DATE_COLUMN], current_date).number_format = "yyyy-mm-dd"
        worksheet.cell(row, columns[COVARIATE_COLUMNS[0]], current_date.isoweekday())
        worksheet.cell(row, columns[COVARIATE_COLUMNS[1]], int(is_workday(current_date)))
        for dish_index, column in enumerate(dishes.values()):
            worksheet.cell(row, column, int(round(max(0.0, float(predicted[dish_index, horizon_index]))))).number_format = "0"
        for column in range(1, worksheet.max_column + 1):
            worksheet.cell(row, column).fill = FORECAST_FILL


def _atomic_save(workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix=".xlsx", delete=False) as temporary:
        temp_path = Path(temporary.name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        temp_path.unlink(missing_ok=True)


def append_forecasts(input_path: Path, output_path: Path, *, days: int = 7, device: str = "cpu") -> ForecastWorkbookReport:
    """Read a source workbook and write a new workbook with blue forecast rows."""
    if days < 1:
        raise ValueError("days 必须大于 0")
    if not input_path.is_file():
        raise FileNotFoundError(f"未找到输入工作簿: {input_path}")
    if input_path.resolve() == output_path.resolve():
        raise ValueError("输出文件必须不同于输入文件，以保护源数据")
    # Keep formulas and styles in the output workbook, but use Excel's cached
    # formula values for the historical target during inference.
    workbook = load_workbook(input_path)
    if workbook.calculation is None:
        workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    else:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    values_workbook = load_workbook(input_path, data_only=True)
    predictor = T0DemandPredictor(device=device)
    rows_written = 0
    for worksheet, values_worksheet in zip(workbook.worksheets, values_workbook.worksheets, strict=True):
        columns = _headers(worksheet)
        dishes = _dish_columns(columns)
        histories, historical_covariates, last_date = _historical_rows(values_worksheet, columns, dishes)
        future_dates, future_covariates = _future_covariates(last_date, days)
        forecast = predictor.predict_many(histories, historical_covariates, future_covariates)
        _append_rows(worksheet, columns, dishes, future_dates, forecast.median)
        rows_written += days
    _atomic_save(workbook, output_path)
    return ForecastWorkbookReport(output_path=output_path, forecast_rows=rows_written)
