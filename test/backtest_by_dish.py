"""Backtest every dish against the final five days of the test workbook."""

from __future__ import annotations

import json
import shutil
import sys
from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from chinese_calendar import is_workday
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from predict.metrics import calculate_metrics
from predict.t0_forecaster import MODEL_ID, T0DemandPredictor


ROOT = Path(__file__).parent
OUTPUT_DIR = ROOT / "预测结果"
CHART_DIR = OUTPUT_DIR / "图表"
TEST_DAYS = 5
DATE_COLUMN = "日期"
COVARIATE_COLUMNS = ("当周的第几天", "是否工作日")
EXCLUDED_COLUMNS = {DATE_COLUMN, *COVARIATE_COLUMNS}


@dataclass(frozen=True)
class DishMetric:
    工作表: str
    菜品: str
    预测天数: int
    MAE: float
    RMSE: float
    MAPE: float
    sMAPE: float


def find_input_workbook() -> Path:
    candidates = [
        path for path in ROOT.glob("*.xlsx") if not path.name.startswith("~$")
    ]
    if len(candidates) != 1:
        raise RuntimeError(f"test 目录应仅有一个输入工作簿，当前找到: {candidates}")
    return candidates[0]


def as_date(value: object) -> date:
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"无效日期: {value!r}")


def prepare_output_directory() -> None:
    if OUTPUT_DIR.exists():
        shutil.rmtree(OUTPUT_DIR)
    CHART_DIR.mkdir(parents=True)


def read_sheet(worksheet) -> tuple[list[date], np.ndarray, dict[str, np.ndarray]]:
    headers = [cell.value for cell in worksheet[1]]
    columns = {
        header: position + 1
        for position, header in enumerate(headers)
        if header is not None
    }
    missing = {DATE_COLUMN, *COVARIATE_COLUMNS}.difference(columns)
    if missing:
        raise ValueError(f"{worksheet.title} 缺少列: {sorted(missing)}")
    dishes = [
        header
        for header in headers
        if header not in EXCLUDED_COLUMNS and header is not None
    ]
    if not dishes:
        raise ValueError(f"{worksheet.title} 未找到菜品列")

    records = []
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row, columns[DATE_COLUMN]).value
        if value is None:
            continue
        current_date = as_date(value)
        day_of_week = worksheet.cell(row, columns[COVARIATE_COLUMNS[0]]).value
        workday = worksheet.cell(row, columns[COVARIATE_COLUMNS[1]]).value
        records.append(
            (
                current_date,
                [
                    float(
                        current_date.isoweekday()
                        if day_of_week is None
                        else day_of_week
                    ),
                    float(
                        int(is_workday(current_date)) if workday is None else workday
                    ),
                ],
                {
                    dish: float(worksheet.cell(row, columns[dish]).value)
                    for dish in dishes
                },
            )
        )
    records.sort(key=lambda record: record[0])
    if len(records) <= TEST_DAYS:
        raise ValueError(f"{worksheet.title} 只有 {len(records)} 天数据，无法完成回测")
    dates = [record[0] for record in records]
    covariates = np.asarray([record[1] for record in records], dtype=np.float32)
    series = {
        dish: np.asarray([record[2][dish] for record in records], dtype=np.float32)
        for dish in dishes
    }
    return dates, covariates, series


def safe_filename(value: str) -> str:
    return "".join(
        "_" if character in '\\/:*?"<>|' else character for character in value
    )


def chart(
    sheet_name: str, dish: str, dates: list[date], actual: np.ndarray, forecast
) -> None:
    plt.figure(figsize=(10, 4.8))
    plt.plot(dates, actual, marker="o", linewidth=2, label="实际销量")
    plt.plot(dates, forecast.median, marker="s", linewidth=2, label="预测销量")
    plt.fill_between(
        dates, forecast.p10, forecast.p90, alpha=0.2, label="10%-90% 预测区间"
    )
    plt.title(f"{sheet_name} - {dish}：逐菜品回测")
    plt.xlabel("日期")
    plt.ylabel("销量")
    plt.grid(alpha=0.25)
    plt.legend()
    plt.tight_layout()
    plt.savefig(
        CHART_DIR / f"{safe_filename(sheet_name)}__{safe_filename(dish)}.png", dpi=160
    )
    plt.close()


def write_workbook(rows: list[dict[str, object]]) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "逐菜品回测明细"
    headers = [
        "工作表",
        "菜品",
        "日期",
        "实际销量",
        "预测销量",
        "P10",
        "P90",
        "误差",
        "绝对误差",
    ]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        worksheet.append([row[header] for header in headers])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column, width in {
        "A": 18,
        "B": 22,
        "C": 14,
        "D": 12,
        "E": 12,
        "F": 10,
        "G": 10,
        "H": 10,
        "I": 12,
    }.items():
        worksheet.column_dimensions[column].width = width
    for row in worksheet.iter_rows(min_row=2, min_col=3, max_col=9):
        row[0].number_format = "yyyy-mm-dd"
        for cell in row[1:]:
            cell.number_format = "0.00"
    output = OUTPUT_DIR / "t0_beta_最后5天回测_逐菜品.xlsx"
    workbook.save(output)
    return output


def main() -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    prepare_output_directory()
    workbook = load_workbook(find_input_workbook(), data_only=True)
    predictor = T0DemandPredictor(device="cpu")
    detail_rows: list[dict[str, object]] = []
    metrics: list[DishMetric] = []

    for worksheet in workbook.worksheets:
        dates, covariates, dishes = read_sheet(worksheet)
        test_start = len(dates) - TEST_DAYS
        history_covariates = covariates[:test_start]
        future_covariates = covariates[test_start:]
        test_dates = dates[test_start:]
        for dish, values in dishes.items():
            forecast = predictor.predict(
                values[:test_start], history_covariates, future_covariates
            )
            actual = values[test_start:]
            predicted = np.maximum(forecast.median, 0)
            metric = calculate_metrics(actual, predicted)
            metrics.append(DishMetric(worksheet.title, dish, TEST_DAYS, **metric))
            chart(worksheet.title, dish, test_dates, actual, forecast)
            for index, current_date in enumerate(test_dates):
                error = float(predicted[index] - actual[index])
                detail_rows.append(
                    {
                        "工作表": worksheet.title,
                        "菜品": dish,
                        "日期": current_date,
                        "实际销量": float(actual[index]),
                        "预测销量": float(predicted[index]),
                        "P10": float(forecast.p10[index]),
                        "P90": float(forecast.p90[index]),
                        "误差": error,
                        "绝对误差": abs(error),
                    }
                )

    output = write_workbook(detail_rows)
    summary = {
        "model": MODEL_ID,
        "train_days": len(dates) - TEST_DAYS,
        "test_days": TEST_DAYS,
        "prediction_granularity": "single_dish",
        "metrics": [asdict(metric) for metric in metrics],
    }
    (OUTPUT_DIR / "metrics.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"已生成 {output}，共回测 {len(metrics)} 个菜品，图表位于 {CHART_DIR}")


if __name__ == "__main__":
    main()
